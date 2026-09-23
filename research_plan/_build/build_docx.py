from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / 'iTE与TG双知识图谱研究计划及后续拓展.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.85), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.05), Cm(2.05)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)

def fonts(style, size, bold=False, east='Songti SC'):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+attr),None)
    rf.set(qn('w:ascii'),'Calibri'); rf.set(qn('w:hAnsi'),'Calibri')
    rf.set(qn('w:eastAsia'),east)
    lang=OxmlElement('w:lang');lang.set(qn('w:eastAsia'),'zh-CN');rpr.append(lang)

fonts(doc.styles['Normal'],11.5)
n=doc.styles['Normal'].paragraph_format
n.line_spacing=Pt(18)
n.space_after=Pt(6)
n.widow_control=True
for name, size in [('Title',21),('Heading 1',16),('Heading 2',12.5)]:
    st=doc.styles[name];fonts(st,size,True,'Heiti SC')
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.line_spacing=1.12
fonts(doc.styles['Subtitle'],11.5,False,'Heiti SC')
doc.styles['Subtitle'].font.italic=False
fonts(doc.styles['List Bullet'],11.5)
doc.styles['List Bullet'].paragraph_format.space_after=Pt(3)

header=sec.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('iTE 与 TG 双知识图谱研究方案');r.font.size=Pt(8)
r.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('第 ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页')

def hyperlink(p,text,url):
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'24516C');rp.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');rp.append(sz)
    r.append(rp);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def table(lines):
    rows=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',x) for x in r)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=[4.0,12.9] if len(rows[0])==2 else [3.3,7.25,6.35]
    if rows[0][0]=='时间':widths=[2.1,8.0,6.8]
    if rows[0][0]=='检查项':widths=[4.0,6.0,6.9]
    if rows[0][0]=='评价维度':widths=[3.2,7.6,6.1]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    props=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
    props.append(b)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=t.rows[i]._tr.get_or_add_trPr()
        split=OxmlElement('w:cantSplit');trpr.append(split)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,(c,txt) in enumerate(zip(cells,row)):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge,value in [('top',90),('bottom',90),('left',110),('right',110)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF4');cp.append(shade)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(15)
            if i==0 or (j==0 and rows[0][0]=='时间'):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(txt);r.font.size=Pt(10.3);r.font.bold=(i==0)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
    p.add_run().font.size=Pt(3)

content=(ROOT/'plan.md').read_text()
lines=content.splitlines();i=0;page=1
while i<len(lines):
    l=lines[i].strip();i+=1
    if not l:continue
    if l=='---PAGE---':
        doc.add_page_break();page+=1;continue
    if l.startswith('|'):
        ts=[l]
        while i<len(lines) and lines[i].strip().startswith('|'):
            ts.append(lines[i].strip());i+=1
        table(ts);continue
    if l.startswith('# '):
        p=doc.add_paragraph(l[2:],'Title');continue
    if l.startswith('## '):
        p=doc.add_paragraph(l[3:],'Heading 1')
        if page>1:p.paragraph_format.space_before=Pt(0)
        continue
    if l.startswith('### '):doc.add_paragraph(l[4:],'Heading 2');continue
    if l.startswith('https://'):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
        hyperlink(p,l,l);continue
    if l.startswith('- '):doc.add_paragraph(l[2:],'List Bullet');continue
    p=doc.add_paragraph(l)
    if l.startswith('版本 '):
        p.paragraph_format.space_after=Pt(12)
        for r in p.runs:r.font.size=Pt(9)
    if l=='研究计划与试验后拓展路线':p.style=doc.styles['Subtitle']
    if page==10:
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=Pt(15.5)
        for r in p.runs:r.font.size=Pt(10.5)

doc.core_properties.title='iTE 与 TG 双知识图谱构建及跨图匹配研究方案'
doc.core_properties.subject='研究计划与试验后拓展路线'
doc.core_properties.author=''
doc.core_properties.keywords='iTE, TG, 知识图谱, 关系匹配, 研究计划'
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)

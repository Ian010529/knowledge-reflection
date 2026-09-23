from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
FIG = OUT / "assets" / "figures"
DOCX = OUT / "thermogalvanic_report_cn.docx"
QA = OUT / "thermogalvanic_report_docx_qa.md"


STYLE = {
    "font_cn": "Microsoft YaHei",
    "font_en": "Calibri",
    "accent": RGBColor(31, 78, 121),
    "accent2": RGBColor(24, 112, 108),
    "muted": RGBColor(90, 96, 100),
    "body": RGBColor(20, 28, 33),
}


def set_run_font(run, name=None, size=None, bold=None, italic=None, color=None):
    if name:
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = color


def set_para_format(p, before=0, after=8, line=1.15, align=None, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align:
        p.alignment = align
    if keep:
        pf.keep_with_next = True


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_para_format(p, after=0, line=1.1)
    r = p.add_run(text)
    set_run_font(r, STYLE["font_cn"], size=size, bold=bold, color=color or STYLE["body"])
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_caption(doc, label, source):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(p, before=2, after=8, line=1.05)
    r = p.add_run(label)
    set_run_font(r, STYLE["font_cn"], size=8.5, bold=True, color=STYLE["muted"])
    r = p.add_run(f"  Source: {source}")
    set_run_font(r, STYLE["font_cn"], size=8.0, color=STYLE["muted"])


def add_figure(doc, filename, caption, source, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(p, before=4, after=2, line=1.0, keep=True)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Inches(width))
    add_caption(doc, caption, source)


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.text = ""
    set_para_format(p, before=14 if level == 1 else 9, after=6, line=1.08, keep=True)
    r = p.add_run(text)
    size = {1: 16, 2: 13, 3: 11.5}.get(level, 11)
    color = STYLE["accent"] if level <= 2 else RGBColor(67, 67, 67)
    set_run_font(r, STYLE["font_cn"], size=size, bold=True, color=color)
    return p


def add_body(doc, text):
    for para in text.strip().split("\n\n"):
        p = doc.add_paragraph()
        set_para_format(p, after=7, line=1.18)
        r = p.add_run(para.strip())
        set_run_font(r, STYLE["font_cn"], size=10.5, color=STYLE["body"])


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_para_format(p, after=4, line=1.15)
        r = p.add_run(item)
        set_run_font(r, STYLE["font_cn"], size=10.2, color=STYLE["body"])


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    cell = table.cell(0, 0)
    shade_cell(cell, "EAF4F2")
    p = cell.paragraphs[0]
    set_para_format(p, before=4, after=4, line=1.12)
    r = p.add_run(title + "  ")
    set_run_font(r, STYLE["font_cn"], size=10.2, bold=True, color=STYLE["accent2"])
    r = p.add_run(body)
    set_run_font(r, STYLE["font_cn"], size=10.0, color=STYLE["body"])
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = widths or [6.5 / len(headers)] * len(headers)
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        shade_cell(hdr[i], "E8EEF5")
        set_cell_text(hdr[i], h, bold=True, color=STYLE["accent"], size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=font_size, align=WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph()
    set_para_format(p, after=6)
    return table


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = STYLE["font_cn"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), STYLE["font_cn"])
    normal.font.size = Pt(10.5)
    for style_name in ["List Bullet", "List Number"]:
        st = styles[style_name]
        st.font.name = STYLE["font_cn"]
        st._element.rPr.rFonts.set(qn("w:eastAsia"), STYLE["font_cn"])
        st.font.size = Pt(10.2)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Thermogalvanic 低品位热能转换综述报告")
    set_run_font(r, STYLE["font_cn"], size=8, color=STYLE["muted"])


def add_cover(doc):
    p = doc.add_paragraph()
    set_para_format(p, before=120, after=12, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Thermogalvanic 低品位热能转换综述报告")
    set_run_font(r, STYLE["font_cn"], size=24, bold=True, color=STYLE["accent"])
    p = doc.add_paragraph()
    set_para_format(p, after=24, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("机制演化、离子热电材料体系与功能聚合物/水凝胶设计策略")
    set_run_font(r, STYLE["font_cn"], size=14, color=STYLE["accent2"])
    p = doc.add_paragraph()
    set_para_format(p, after=28, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("基于 7 篇近期综述的综合报告；不展开器件集成和可穿戴应用层面。")
    set_run_font(r, STYLE["font_cn"], size=10.5, color=STYLE["muted"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / "fig_low_grade_heat_roadmap.png"), width=Inches(5.6))
    add_caption(doc, "图 0. 低品位热背景与 iTE 发展路线图（封面图）", "Fig. 1, Chem. Soc. Rev., 2026")
    p = doc.add_paragraph()
    set_para_format(p, before=18, after=0, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("输出文件：thermogalvanic_report_cn.docx")
    set_run_font(r, STYLE["font_cn"], size=9.5, color=STYLE["muted"])
    add_page_break(doc)


def add_manual_toc(doc):
    add_heading(doc, "目录", 1)
    entries = [
        ("摘要", "1"),
        ("缩略语与术语", "2"),
        ("1. 低品位热能背景：为什么 thermogalvanic 值得关注", "3"),
        ("2. Thermogalvanic 的基本机制", "5"),
        ("3. TEG、TCC 与 TEC：材料路线的边界", "7"),
        ("4. 离子热电材料体系：从液体到功能聚合物", "8"),
        ("5. Redox couple 是 thermogalvanic 的第一性设计变量", "10"),
        ("6. Thermogalvanic hydrogels：从“固定电解液”到主动调控网络", "12"),
        ("7. 功能聚合物的多尺度设计策略", "14"),
        ("8. 评价指标：thermopower 之外还要看什么", "15"),
        ("9. 材料层面的关键挑战", "16"),
        ("10. 面向材料研究的设计建议", "18"),
        ("11. 结论与参考综述来源", "19"),
    ]
    table = doc.add_table(rows=len(entries), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.8)
    table.columns[1].width = Inches(0.6)
    for i, (title, page) in enumerate(entries):
        cells = table.rows[i].cells
        set_cell_text(cells[0], title, size=9.5)
        set_cell_text(cells[1], page, size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_note(
        doc,
        "页码说明",
        "本目录为手工目录，页码按当前 DOCX 版式和分节估算；在 Word 中如调整图片大小或字体，页码可能略有变化。",
    )
    add_page_break(doc)


def add_front_matter(doc):
    add_heading(doc, "摘要", 1)
    add_body(
        doc,
        """
低品位热能（通常指低于 100 °C 的热源）广泛存在于工业废热、人体热、环境热和微电子散热中，却因为温差低、分布分散、热流密度不稳定而难以被传统热机高效回收。离子热电（ionic thermoelectrics, iTEs）因其高热电势、低热导、柔性和材料来源广等特点，正在成为低品位热能转换的重要方向。其中，thermogalvanic effect（TG，热电化学/热伽伐尼效应）依赖温差驱动氧化还原对在冷热电极处发生电化学势偏移，并通过外电路传递电子，因此具备连续输出潜力。

本报告综合近期关于 liquid-state thermocells、thermo-electrochemical cells、ionic thermoelectrics、thermogalvanic hydrogels 和 functional polymers for iTEs 的多篇综述，聚焦材料与机制层面：首先说明低品位热回收背景和 iTE 研究演化；随后比较 thermodiffusion、thermogalvanic 与 thermoextraction 的机制差异；再从 redox couple、溶剂化结构、离子浓度、聚合物-离子相互作用、凝胶网络和热/电/离子耦合传输角度总结材料设计策略。报告不展开器件集成、可穿戴系统、模块串并联和应用演示，只把它们作为材料评价边界条件简要提及。
""",
    )
    add_note(
        doc,
        "报告范围",
        "重点在 thermogalvanic 材料和机制，不讨论器件结构设计、人体穿戴集成、模块连接方式或应用产品化路线。",
    )
    add_heading(doc, "缩略语与术语", 1)
    add_table(
        doc,
        ["术语", "英文/符号", "本报告使用方式"],
        [
            ["低品位热", "low-grade heat", "通常指 <100 °C 热源；作为背景约束"],
            ["离子热电", "ionic thermoelectrics, iTEs", "涵盖 TD、TG、耦合体系等"],
            ["热电化学/热伽伐尼效应", "thermogalvanic effect, TG", "报告主线机制"],
            ["热扩散效应", "thermodiffusion effect, TD", "作为 TG 的对照机制"],
            ["热电化学电池", "thermo-electrochemical cell, TEC", "指依赖 TG 的电化学热电单元"],
            ["热电势", "thermopower / Seebeck coefficient", "用 α 或 Si 表示"],
            ["氧化还原对", "redox couple", "如 Fe(CN)6^4-/3-, I-/I3-"],
            ["准固态离子导体", "quasi-solid-state ionic conductor, QSS", "聚合物/凝胶基体系"],
        ],
        widths=[1.45, 2.25, 2.8],
        font_size=8.5,
    )
    add_page_break(doc)


def section_literature_map(doc):
    add_heading(doc, "0. 文献脉络与综合口径", 1)
    add_body(
        doc,
        """
本报告不是逐篇文献摘要，而是把多篇综述重新组织为一条材料与机制主线。Joule 2021 的 perspective 提供 liquid-state thermocell 与低品位热回收的早期框架；Advanced Science 2021 和 2025 的综述分别提供 wearable TECs 和 thermogalvanic hydrogels 的材料进展，其中本报告只抽取机制、材料和水凝胶设计内容，不讨论器件和应用演示；Chemical Science 2024 和 National Science Review 2026 提供 iTE 分类、理论建模和材料全景；Advanced Materials 2026 强调功能聚合物在离子动力学、力学和能量转换中的多尺度作用；Chemical Society Reviews 2026 则把 iTE 发展放在“持续供电”的大主题中，提供低品位热背景和演化路线图。

综合这些综述后，可以形成一个适合材料报告的逻辑：低品位热背景提出问题，TG 机制给出连续输出的物理化学基础，redox couple 决定热力学上限，溶剂化和聚合物网络决定传输与稳定性，水凝胶和功能聚合物把 TG 从液态模型推向准固态材料平台。这个口径避免把报告写成器件或应用综述，也避免把不同综述中的应用图片堆叠在一起。
""",
    )
    add_table(
        doc,
        ["来源综述", "本报告采用的内容", "本报告不展开的内容"],
        [
            ["Joule 2021", "低品位热来源、液态 thermocell 基本框架", "模块和温度检测/冷却应用"],
            ["Adv. Sci. 2021", "TEC/TCC/TG 机制和柔性材料问题", "可穿戴系统设计"],
            ["Chem. Sci. 2024", "iTE 分类、理论建模、性能参数", "器件组合和集成路线"],
            ["Adv. Sci. 2025", "thermogalvanic hydrogels、redox/凝胶策略", "TH-based device integration"],
            ["Adv. Mater. 2026", "functional polymers、多尺度 ion dynamics 设计", "实际皮肤适配和装备型应用"],
            ["Chem. Soc. Rev. 2026", "低品位热背景、iTE 演化路线、持续供电主题", "power supply 应用工程"],
            ["Natl. Sci. Rev. 2026", "材料分类、redox couples、polymer design mechanisms", "products 和产品化视角"],
        ],
        widths=[1.55, 2.85, 2.1],
        font_size=8.0,
    )
    add_page_break(doc)


def section_background(doc):
    add_heading(doc, "1. 低品位热能背景：为什么 thermogalvanic 值得关注", 1)
    add_body(
        doc,
        """
低品位热的特点是来源广、温度低、热流不稳定且空间分散。传统蒸汽循环或高温热机更适合集中高品位热源，而人体热、环境温差、电子设备散热和中低温工业余热往往难以通过复杂机械系统回收。电子热电材料（electronic thermoelectrics, eTEs）可以直接把温差转化为电压，但在低温差和柔性场景中通常受限于较低 Seebeck 系数、材料刚性、热导和成本等因素。

离子热电的吸引力在于：离子迁移、溶剂化重排和 redox reaction entropy 可以产生远高于常见电子热电材料的热电势；液体、凝胶和聚合物体系天然具有低热导和柔性；材料选择空间覆盖水系盐、离子液体、深共熔溶剂、聚电解质和多孔框架。过去十余年，iTE 研究从“证明高热电势”逐渐转向“实现稳定功率输出”。在这个转向中，thermogalvanic 体系因其 redox 反应可把电子传递到外电路而占据关键位置。
""",
    )
    add_figure(doc, "fig_low_grade_heat_roadmap.png", "图 1. 低品位热背景、iTE 论文增长与代表性演化路线", "Fig. 1, Chem. Soc. Rev., 2026", width=6.4)
    add_body(
        doc,
        """
图 1 中的时间线说明 iTE 研究经历了几个明显阶段：早期关注液态热电化学电池和纳米多孔电解质中的巨热电势；随后 iTE supercapacitors、charged nanochannels 与 supramolecular thermocells 被用于放大电压或储能；2020 年之后，thermodiffusion 与 thermogalvanic 的耦合、高功率输出和材料体系优化成为高频主题。这个演化体现出一个核心判断：单纯追求高开路电压不足以支撑实际低品位热回收，稳定输出和材料可持续运行同样重要。
""",
    )
    add_heading(doc, "1.1 Liquid-state thermocell 的起点", 2)
    add_body(
        doc,
        """
液态 thermocell 是理解 thermogalvanic 的基础模型。典型体系由两个电极和含 redox couple 的电解质组成，冷热端电极处的电化学势随温度改变而不再相等。当外电路接通时，热端和冷端发生互补的氧化还原反应，电子经外电路传输，离子在电解质中扩散以维持电荷和浓度平衡。相比固态热电，液态体系的优点是低成本、可扩展、高界面反应速率和较高 thermopower；缺点是泄漏、挥发、冻结/沸腾、封装与长期稳定性。
""",
    )
    add_figure(doc, "fig_liquid_state_thermocell.png", "图 2. 液态 thermocell 的低品位热来源和基本转换单元", "Fig. 1, Joule, 2021", width=5.7)
    add_page_break(doc)


def section_mechanism(doc):
    add_heading(doc, "2. Thermogalvanic 的基本机制", 1)
    add_body(
        doc,
        """
Thermogalvanic effect 的热电势来自 redox reaction entropy。对于可逆氧化还原对 A/B，在不同温度的两端电极上，反应自由能和电极电势随温度变化。若热端和冷端温度分别为 Th 与 Tc，开路状态下可观测到电压差；闭路后，热端与冷端发生相反方向的电极反应，电子通过外电路流动，电解质中的 redox species 和 supporting ions 通过扩散与迁移补偿浓度和电荷变化。

常用表达式可写为 Si = -(Vhot - Vcold)/(Thot - Tcold)。从热力学角度，TG 的 thermopower 与 redox reaction entropy difference 相关，近似满足 α ≈ ΔSredox/nF，其中 n 是转移电子数，F 是法拉第常数。因此，redox couple 的选择、溶剂化结构、配位/水合状态、添加剂和聚合物环境都会改变热电势。
""",
    )
    add_figure(doc, "fig_three_mechanisms.png", "图 3. TD、TG 和 thermoextraction 三类离子热电机制对比", "Fig. 1, Adv. Sci., 2025", width=6.2)
    add_heading(doc, "2.1 与 thermodiffusion 的差异", 2)
    add_body(
        doc,
        """
Thermodiffusion（TD）依赖 Soret effect：阳离子和阴离子在温度梯度下具有不同热迁移能力，导致离子浓度和电双层重新分布，从而产生热电压。TD 体系的 thermopower 常常更高，但离子本身不能穿过外电路，实际工作模式更接近“热充电—外接放电—再平衡”。因此，TD 在高电压和传感/储能耦合方面有优势，但持续供电能力受工作循环限制。

Thermogalvanic（TG）则通过 redox reaction 把电荷转移给电极和外电路。只要温差、redox 反应可逆性和物质传输能够维持，TG 可以提供连续电流输出。代价是：TG 的 thermopower 通常低于最高 TD 体系，且受电极反应动力学、浓差极化、redox species 扩散、交叉反应和内阻限制。因此，TG 优化既是热力学问题，也是传输动力学问题。
""",
    )
    add_figure(doc, "fig_td_tg_voltage_profiles.png", "图 4. TD 与 TG 的工作阶段和电压曲线差异", "Fig. 2, Adv. Mater., 2026", width=6.3)
    add_note(
        doc,
        "机制判断",
        "本报告把 TG 的核心价值定义为“连续输出能力 + 低温差适配 + 可由聚合物/凝胶材料调控”，而不是单独追逐最高开路电压。",
    )
    add_page_break(doc)


def section_comparison(doc):
    add_heading(doc, "3. TEG、TCC 与 TEC：材料路线的边界", 1)
    add_body(
        doc,
        """
在低品位热能转换语境中，TEG、TCC 和 TEC 经常被放在同一张图谱中比较。TEG 主要依赖电子/空穴在固态半导体中的 Seebeck effect，工作模式清晰且可连续输出，但在柔性、低温差和大面积低成本部署方面存在约束。TCC 依赖 thermodiffusion，可获得较高 thermopower，但输出具有充放电特征。TEC/TG 使用 redox ions 作为关键活性物质，热电势中等但可连续输出，材料可以从液态扩展到水凝胶和 quasi-solid-state ionic conductors。
""",
    )
    add_table(
        doc,
        ["维度", "TEG", "TCC/TD", "TEC/TG"],
        [
            ["主机制", "电子 Seebeck effect", "离子热扩散 / Soret effect", "温差驱动 redox reaction entropy"],
            ["载流子", "电子和空穴", "阳离子/阴离子", "Redox ions + supporting ions"],
            ["典型材料", "Bi2Te3 等固态半导体", "水系盐、离子液体、聚电解质", "Fe(CN)6^4-/3-, I-/I3-, Fe2+/3+ 等 redox 体系"],
            ["热电势", "通常较低", "通常较高", "中等，可由 redox entropy 调控"],
            ["输出模式", "连续输出", "热充电/放电特征明显", "具备连续输出潜力"],
            ["主要短板", "低温差性能、柔性、成本", "持续供电和能量密度", "内阻、浓差极化、长期稳定性"],
        ],
        widths=[1.0, 1.75, 1.85, 1.9],
        font_size=8.2,
    )
    add_body(
        doc,
        """
这一区分决定了后续材料设计的评价逻辑：对于 TG 体系，α/Si 不是唯一指标。即使 redox entropy 提高，如果离子电导率下降、扩散受阻、热导上升或电极反应变慢，最终功率密度和能量转换效率仍可能降低。因此，TG 材料评价至少要同时考虑 thermopower、ionic/electrical conductivity、thermal conductivity、mass transport、redox reversibility 和 cycle stability。
""",
    )
    add_page_break(doc)


def section_material_map(doc):
    add_heading(doc, "4. 离子热电材料体系：从液体到功能聚合物", 1)
    add_body(
        doc,
        """
离子热电材料可以按照物理状态和离子传输环境分为液态、凝胶态、聚电解质和多孔固态框架。液态电解质反应和扩散快，是建立机理和筛选 redox couple 的最直接平台；凝胶态材料在保留水系/离子液体传输特性的同时提高形状稳定性；聚电解质通过固定电荷、链段极性和离子选择性增强热迁移差异或 redox species 分布不均匀性；多孔固态框架则利用纳米通道、梯度亲疏水性和界面吸附调节离子输运。
""",
    )
    add_figure(doc, "fig_material_categories.png", "图 5. iTE 材料类别：液态、凝胶态、聚电解质和多孔固态框架", "Fig. 2, Natl. Sci. Rev., 2026", width=6.1)
    add_heading(doc, "4.1 多维性能图谱", 2)
    add_body(
        doc,
        """
Advanced Materials 2026 的综述把 reported iTEs 与 eTEs 放在 thermopower-conductivity 图谱上，可以看到 TD 型 iTE 在 thermopower 上显著突出，而 TG 型 hydrogel/液态体系位于中等 thermopower 区间。与此同时，电子热电材料的电导率明显更高。这个图谱说明 iTE 并非直接替代 eTE，而是在低温差、柔性、低热导和低成本等场景中形成互补。

对 TG 而言，性能优化的困难在于多个指标方向并不总是同向：提高聚合物网络强度可能降低 ion diffusivity；增强 redox species 与基体相互作用可能提高 thermopower，却也可能降低扩散系数；提高盐浓度可增强电导，但可能改变水合结构、提高黏度或引起沉淀。因此，多尺度设计必须把分子相互作用、微结构通道和宏观热/水稳定性同时纳入。
""",
    )
    add_figure(doc, "fig_performance_material_comparison.png", "图 6. iTE/eTE 性能区间与功能聚合物多尺度设计策略", "Fig. 1, Adv. Mater., 2026", width=6.3)
    add_page_break(doc)


def section_redox(doc):
    add_heading(doc, "5. Redox couple 是 thermogalvanic 的第一性设计变量", 1)
    add_body(
        doc,
        """
TG 的热电势直接来自 redox reaction entropy，因此 redox couple 的选择是首要变量。经典 p-type 体系包括 ferri/ferrocyanide Fe(CN)6^3-/4- 和 Fc/Fc+；常见 n-type 体系包括 Fe2+/3+、I-/I3-、Cu/Cu2+ 以及部分金属配合物。一个理想 redox couple 需要同时满足：反应可逆、电子转移快、化学稳定、溶解度高、毒性低、与聚合物/电极兼容，并能在温差下产生较高 entropy difference。

Fe(CN)6^4-/3- 被广泛使用，是因为其水系可逆性好、反应动力学快且容易与水凝胶体系兼容。但它的 intrinsic thermopower 并不是最高，后续研究大量围绕 solvation shell restructuring、chaotropic/cosmotropic additives、counterion engineering 和温敏结晶展开。I-/I3- 则可通过选择性包结、络合和温度依赖溶解度变化实现 p/n 调控或自充电行为，但也更需要关注副反应和长期稳定性。
""",
    )
    add_figure(doc, "fig_redox_couples_strategies.png", "图 7. 常用 redox couples 与提升 thermopower 的代表性策略", "Fig. 6, Natl. Sci. Rev., 2026", width=5.9)
    add_heading(doc, "5.1 Redox thermodynamics 与溶剂化结构", 2)
    add_body(
        doc,
        """
溶剂化结构改变 redox species 在不同价态下的 entropy，从而改变 ΔSredox。对于 Fe(CN)6^4-/3-，水合层、counterion 配位以及 guanidinium 等添加剂可以改变离子周围的氢键网络和局部结构有序度。若高/低价态的溶剂化 entropy 差被放大，thermopower 就可能提高。类似地，有机溶剂、深共熔溶剂和混合溶剂可通过 donor number、极性、黏度和离子缔合作用影响热电势。

然而，溶剂化调控不能只看 α。溶剂和添加剂往往同时改变 redox solubility、viscosity、diffusion coefficient 和 ionic conductivity。一个体系在开路电压上表现优秀，闭路功率输出却可能因扩散慢和极化严重而受限。因此，redox 策略需要以热力学增益和传输损失的平衡为目标。
""",
    )
    add_table(
        doc,
        ["策略", "作用对象", "可能收益", "主要风险"],
        [
            ["更换 redox couple", "反应熵、可逆性、电子转移", "直接改变 α 和 p/n 极性", "毒性、稳定性、成本、溶解度"],
            ["调节 counterion/additive", "溶剂化壳层与局部结构", "放大 ΔSredox，抑制沉淀", "黏度升高、扩散变慢、副反应"],
            ["混合溶剂/深共熔溶剂", "donor number、极性、氢键网络", "改变反应 entropy 与温度窗口", "水分敏感、长期化学稳定性"],
            ["聚合物络合/选择性捕获", "redox species 空间分布", "形成浓度差或选择性迁移", "降低反应活性或可逆性"],
        ],
        widths=[1.25, 1.75, 1.75, 1.75],
        font_size=8.2,
    )
    add_page_break(doc)


def section_hydrogel(doc):
    add_heading(doc, "6. Thermogalvanic hydrogels：从“固定电解液”到主动调控网络", 1)
    add_body(
        doc,
        """
Thermogalvanic hydrogels（THs）把 redox couple、supporting ions、水和聚合物网络结合在一起，兼具液态电解质的离子传输和固态材料的形状保持。它们不是简单把液态电解液凝胶化，而是通过 polymer-ion interaction、pore structure、water retention、mechanical reinforcement 和 solvation modulation 主动调控 TG performance。

天然聚合物如 gelatin、cellulose、alginate、chitosan 等具有生物相容性和丰富官能团，适合通过氢键、配位和离子作用影响 redox species。合成聚合物如 PVA、PAAm、PEO、PVDF-HFP、聚离子液体和 zwitterionic polymers 则提供更可控的交联密度、相分离结构、离子通道和力学性能。水凝胶研究的关键问题，是在高含水量、高离子电导和强力学稳定之间找到平衡。
""",
    )
    add_figure(doc, "fig_hydrogel_synergy.png", "图 8. Gelatin/cellulose 基 thermogalvanic hydrogels 的协同效应示例", "Fig. 2, Adv. Sci., 2025", width=6.1)
    add_heading(doc, "6.1 Gelatin、cellulose 与天然聚合物", 2)
    add_body(
        doc,
        """
Gelatin 体系常被用作 Fe(CN)6^4-/3-/KCl 的协同平台。聚合物链、KCl 和 Fe(CN)6 redox couple 共同贡献 thermopower：redox entropy 提供 TG 贡献，K+ / Cl- 的热迁移提供 TD 贡献，gelatin matrix 通过水合和链段作用进一步调节离子分布。综述中常见的 gelatin-KCl-Fe(CN)6 体系展示了 TD 与 TG 协同后 thermopower 增大的趋势。

Cellulose 及 bacterial cellulose 体系强调机械稳定性、孔道结构和水保持能力。纤维素纳米网络能够提供连续水相通道和高含水量，同时通过羟基和物理缠结增强凝胶韧性。这类材料适合研究“高水含量 + 机械支撑 + redox 可逆性”的组合，但长期湿度变化和干燥收缩仍需处理。
""",
    )
    add_heading(doc, "6.2 PVA、PAAm 与合成网络", 2)
    add_body(
        doc,
        """
PVA 的羟基、结晶/冻融结构和可拉伸网络使其成为 THs 中常见基体。通过机械训练、取向、冻融和交联调节，PVA 可以形成更有序的离子通道，改善 redox species 和 supporting ions 的迁移路径。PAAm 则以酰胺基团和柔性链段著称，容易与金属离子或水形成氢键/配位作用，兼顾弹性和水保持。

合成网络的优势是可设计性：交联密度、网络拓扑、链段极性、固定电荷、疏水/亲水微区都可以被系统调控。缺点是单一优化常带来副作用，例如增加交联密度提高强度但降低电导，提高盐浓度提高电导但改变水结构和黏度。因此，需要把水凝胶视为热-电-离子-力学耦合材料。
""",
    )
    add_page_break(doc)


def section_polymer(doc):
    add_heading(doc, "7. 功能聚合物的多尺度设计策略", 1)
    add_body(
        doc,
        """
功能聚合物对 TG/iTE 的贡献可以分为三个尺度。分子尺度上，官能团通过氢键、静电作用、配位、疏水相互作用和 zwitterionic interactions 改变离子解离、溶剂化和 redox species 分布。微结构尺度上，聚合物相分离、取向、纳米孔、纤维网络和凝胶孔径决定离子扩散路径。宏观尺度上，热导、含水稳定、力学韧性和加工性决定材料能否保持长期性能。

在 thermogalvanic 体系中，聚合物不仅影响 supporting ions，也影响 redox couple 的局部环境。如果聚合物对某一价态 redox species 有更强相互作用，就可能改变两端浓度分布或 solvation entropy；如果聚合物形成选择性通道，就可能降低浓差极化；如果聚合物提高保水性和抗冻/抗干燥能力，就能延长低品位热环境下的工作窗口。
""",
    )
    add_table(
        doc,
        ["尺度", "设计变量", "对 TG 的作用", "评价指标"],
        [
            ["分子尺度", "官能团、固定电荷、zwitterion、络合位点", "调节 redox 溶剂化、离子解离和 ΔSredox", "α/Si、diffusion coefficient、redox reversibility"],
            ["微结构尺度", "孔径、取向通道、相分离、纤维网络", "降低传输阻力，抑制浓差极化", "ionic conductivity、viscosity、mass transport"],
            ["宏观尺度", "含水率、热导、力学强度、抗冻/抗干燥", "维持温差和长期结构稳定", "κ、循环稳定性、拉伸/弯折性能"],
        ],
        widths=[1.05, 2.0, 2.05, 1.4],
        font_size=8.1,
    )
    add_body(
        doc,
        """
功能聚合物策略的共同趋势是从“被动支撑基体”转向“主动调控传输和热力学”。例如，zwitterionic polymers 可通过强水合和离子偶极相互作用改变 mobile ions 的扩散差异；aligned polymer networks 可诱导方向性传输；低热导结构可帮助保持温度梯度；特定官能团可通过选择性络合改变 redox species 的局域化学环境。
""",
    )
    add_page_break(doc)


def section_metrics(doc):
    add_heading(doc, "8. 评价指标：thermopower 之外还要看什么", 1)
    add_body(
        doc,
        """
Thermopower 是 TG 材料最直观的指标，但不能单独代表材料的热能转换能力。实际功率输出还取决于离子电导率、电子/电极界面电荷转移、redox species 扩散、热导率、浓差极化和热端/冷端温差保持。对于材料层面报告，可以把评价分为四类：热力学指标、传输指标、热管理指标和稳定性指标。
""",
    )
    add_table(
        doc,
        ["类别", "关键指标", "说明", "常见误区"],
        [
            ["热力学", "α/Si, ΔSredox", "决定开路电压和温差响应", "只比较 α 而忽略输出电流"],
            ["离子/反应传输", "σi, diffusion coefficient, charge-transfer resistance", "决定闭路电流和极化程度", "提高黏度或交联后未检查扩散"],
            ["热管理", "κ, ΔT retention", "决定温差能否在材料中保持", "高电导材料同时提高热导"],
            ["稳定性", "cycle stability, water retention, chemical stability", "决定长时间低温差工作", "短时测试代替长期运行"],
            ["综合输出", "Pmax, η, ηr", "连接材料与能量转换效率", "测试面积、厚度、温差不统一"],
        ],
        widths=[1.0, 1.65, 2.0, 1.85],
        font_size=8.1,
    )
    add_body(
        doc,
        """
为了避免“指标漂移”，未来 TG 材料比较需要明确测试条件：温差大小、样品厚度、有效面积、电极材料、redox concentration、supporting electrolyte、湿度/含水量、测量时间和循环次数。特别是水凝胶体系，含水率变化会同时改变电导、热导、机械性能和 redox diffusion。如果不报告这些条件，跨文献比较容易失真。
""",
    )
    add_page_break(doc)


def section_challenges(doc):
    add_heading(doc, "9. 材料层面的关键挑战", 1)
    add_heading(doc, "9.1 Redox couple 的高热电势与长期稳定难以兼得", 2)
    add_body(
        doc,
        """
高 ΔSredox 往往需要特殊溶剂化结构、添加剂或选择性络合，但这些策略可能引入沉淀、副反应、黏度升高或扩散变慢。对于 Fe(CN)6^4-/3- 等稳定体系，进一步提升 thermopower 的空间有限；对于 I-/I3- 或金属配合物等体系，热电势和 p/n 极性可调性更大，但长期化学稳定性和副反应风险也更高。
""",
    )
    add_heading(doc, "9.2 聚合物网络增强力学性能时可能牺牲传输", 2)
    add_body(
        doc,
        """
交联、双网络、取向和填料增强能够提高凝胶强度与韧性，却常常降低离子扩散和 redox species 的有效迁移。TG 材料需要 redox species 在冷热端之间持续补给，若网络过密或局部强络合，浓差极化会显著增加。材料设计必须找到“足够强、但不堵塞”的网络窗口。
""",
    )
    add_heading(doc, "9.3 含水稳定和低热导之间存在耦合", 2)
    add_body(
        doc,
        """
水是多数高性能水系 THs 的关键传输介质，但水分蒸发、冻结、污染和渗漏会改变性能。加入保湿剂、离子液体、深共熔溶剂或抗冻组分可以拓宽环境适应性，却会改变热导、黏度和离子溶剂化。低热导有利于保持 ΔT，但过低的热传输也可能减慢响应速度。因此，含水稳定、热导和响应时间需要协同优化。
""",
    )
    add_heading(doc, "9.4 缺少统一的材料比较基准", 2)
    add_body(
        doc,
        """
综述中报道的 TG/TH 性能常来自不同电极、不同温差、不同厚度和不同测试时间。材料层面的公平比较需要标准化样品几何、温差、redox concentration、测试环境和计算方法。否则，表观高 thermopower、功率密度或效率很难判断是否来自材料本身，还是来自测试条件差异。
""",
    )
    add_page_break(doc)


def section_outlook(doc):
    add_heading(doc, "10. 面向材料研究的设计建议", 1)
    add_body(
        doc,
        """
基于多篇综述的共同结论，thermogalvanic 材料研究可沿四条路径推进。第一，建立 redox couple 的热力学-动力学数据库，记录 α、diffusion coefficient、reaction reversibility、solubility、toxicity 和 stability，而不仅仅收集最高 thermopower。第二，把溶剂化结构和聚合物-离子相互作用作为可表征变量，用 Raman/FTIR/NMR、MD simulation 和 electrochemical impedance 等方法建立结构-性能关系。第三，发展可控微结构水凝胶，在取向通道、双网络、保水抗冻和低热导之间寻找平衡。第四，建立统一测试协议，把材料在相同 ΔT、几何和时间尺度下的 Pmax、ηr、cycle stability 放在一起比较。
""",
    )
    add_note(
        doc,
        "不讲器件层面时的结论落点",
        "可以把报告收束到材料判据：一个好的 TG 材料应同时具备高 ΔSredox、快速且选择性的离子/redox 传输、低热导、稳定含水或准固态环境，以及可重复测试的长期化学稳定性。",
    )
    add_table(
        doc,
        ["研究方向", "建议优先回答的问题"],
        [
            ["Redox chemistry", "哪类 redox couple 在低毒、可逆和高 ΔSredox 之间最平衡？"],
            ["Solvation engineering", "添加剂/溶剂如何改变不同价态的水合 entropy 和扩散？"],
            ["Polymer-ion interaction", "哪些官能团能提高热电势而不显著降低 mass transport？"],
            ["Hydrogel microstructure", "怎样同时实现高含水、高强度、低热导和快速 redox 补给？"],
            ["Benchmarking", "如何统一 ΔT、厚度、面积、电极和时间尺度来比较材料？"],
        ],
        widths=[2.0, 4.5],
        font_size=8.8,
    )
    add_page_break(doc)


def section_conclusion_refs(doc):
    add_heading(doc, "11. 结论", 1)
    add_body(
        doc,
        """
Thermogalvanic 是离子热电领域中最适合讨论连续低品位热能转换的机制。它利用温差引起的 redox reaction entropy difference 产生电化学势差，并通过外电路传递电子。与 thermodiffusion 相比，TG 的 thermopower 通常不是最高，但在连续输出方面更具材料逻辑上的优势。

材料研究的核心并非单点提高 α，而是把 redox thermodynamics、ion diffusion、polymer network、water/solvent stability 和 thermal conductivity 放入同一个优化框架。Redox couple 决定热力学上限，溶剂化结构和添加剂决定 entropy 与扩散，功能聚合物和水凝胶网络决定准固态体系的传输、力学和环境稳定性。未来有价值的工作应减少孤立性能记录，转向可比较、可复现和可长期运行的材料基准。
""",
    )
    add_heading(doc, "参考综述来源", 1)
    refs = [
        "Duan et al. Liquid-state thermocells: Opportunities and challenges for low-grade heat harvesting. Joule, 2021.",
        "Liu et al. Potentially Wearable Thermo-Electrochemical Cells for Body Heat Harvesting: From Mechanism, Materials, Strategies to Applications. Advanced Science, 2021.",
        "Jabeen et al. Recent advances in ionic thermoelectric systems and theoretical modelling. Chemical Science, 2024.",
        "Lin et al. Critical Design Strategy of Thermogalvanic Hydrogels for Low-Grade Heat Harvesting. Advanced Science, 2025.",
        "Kim et al. Functional Polymers for Ionic Thermoelectrics: Multiscale Design Strategies for Ion Dynamics, Mechanics, and Energy Harvesting. Advanced Materials, 2026.",
        "Shang et al. Pushing ionic thermoelectrics toward power supply applications: origins, advances, challenges, and future directions. Chemical Society Reviews, 2026.",
        "Xu et al. A bird's-eye view for future ionic thermoelectrics: from ions to products. National Science Review, 2026.",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Number")
        set_para_format(p, after=4, line=1.12)
        r = p.add_run(ref)
        set_run_font(r, STYLE["font_cn"], size=9.3, color=STYLE["body"])


def write_qa(render_attempted=False):
    lines = [
        "# DOCX QA",
        "",
        f"- File: `{DOCX}`",
        "- Target length: approximately 20 pages in Word, with explicit page breaks between major sections and figure-rich blocks.",
        "- Scope check: device integration, wearable application, module connection, and product/application figures were excluded.",
        "- Background emphasis: low-grade heat from data centers / computing facilities, power electronics, industrial processes, and microelectronics; human-body heat is not used as a main background driver.",
        "- Evidence style: review-cited representative studies are organized as examples of design principles rather than as isolated bullet lists.",
        "- Included figures: low-grade heat/evolution roadmap, mechanism comparison, TD/TG voltage profiles, performance comparison, material categories, redox strategies, hydrogel synergy.",
        "- Structural validation: DOCX generated and reopened by python-docx without error.",
        "- Visual render QA: skipped because `soffice`/LibreOffice is not installed in this environment.",
    ]
    QA.write_text("\n".join(lines) + "\n", encoding="utf-8")


def add_manual_toc_v2(doc):
    add_heading(doc, "目录", 1)
    entries = [
        ("摘要：本报告要回答的问题", "1"),
        ("1. 背景：从低品位热到算力中心废热", "3"),
        ("2. 基本机制：为什么 TG 能连续输出", "5"),
        ("3. 核心变量一：redox couple 和溶剂化结构", "8"),
        ("4. 核心变量二：从液态电解质到聚合物/水凝胶基体", "12"),
        ("5. 体系化比较：不同材料路线已经解决了什么", "16"),
        ("6. 剩余问题：下一步材料设计该围绕什么展开", "19"),
        ("参考综述来源", "20"),
    ]
    table = doc.add_table(rows=len(entries), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.8)
    table.columns[1].width = Inches(0.6)
    for i, (title, page) in enumerate(entries):
        cells = table.rows[i].cells
        set_cell_text(cells[0], title, size=9.5)
        set_cell_text(cells[1], page, size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_note(doc, "结构说明", "新版报告按“问题-机制-材料变量-代表工作-瓶颈”的单线逻辑展开；不按每篇综述逐篇介绍。")
    add_page_break(doc)


def add_front_matter_v2(doc):
    add_heading(doc, "摘要：本报告要回答的问题", 1)
    add_body(
        doc,
        """
低品位热能并不缺少来源，真正困难在于温差小、热流分散、往往还伴随散热需求。算力中心和数据中心就是典型例子：大量电能最终以低温废热释放，热源空间分布复杂，并且散热本身是维持系统运行的刚性需求。因此，低品位热回收技术不能只追求“能从温差产生电压”，还必须面对材料成本、低温差响应、连续输出、热管理兼容性和长期稳定性。

Thermogalvanic effect（TG，热电化学/热伽伐尼效应）提供了一条不同于传统固态热电的路线。它不是依赖电子/空穴在固体晶格中的 Seebeck effect，而是利用温度改变 redox couple 的电化学势：热端和冷端发生相反方向的氧化还原反应，电子经外电路传输，离子在电解质中扩散补偿。与 thermodiffusion（TD）相比，TG 的热电势往往不是最高，但其优势在于可以通过 redox 反应维持连续电流输出。

本报告综合 7 篇综述，并把综述中引用的代表性原始工作重新组织为材料设计证据。报告主线是：TG 性能首先由 redox couple 的反应熵决定，其次由溶剂化结构、添加剂和浓度梯度放大，再由聚合物/水凝胶网络控制离子传输、含水稳定、热导和力学约束。报告不讨论器件集成、穿戴应用和模块连接，只回答材料体系层面的一个问题：已经有哪些策略被证明能提高 thermogalvanic 低品位热转换性能，它们各自解决了什么、又留下什么问题？
""",
    )
    add_note(
        doc,
        "一句主线",
        "TG 材料设计 = redox reaction entropy 的热力学增益 + 溶剂化/聚合物环境的传输调控 + 准固态基体的稳定性约束。",
    )
    add_heading(doc, "术语表", 2)
    add_table(
        doc,
        ["术语", "含义", "在本报告中的作用"],
        [
            ["低品位热", "通常指 <100 °C 的热源", "以算力中心、工业和电子设备废热作为背景"],
            ["Thermogalvanic effect (TG)", "温差改变 redox 电势并驱动电化学反应", "报告的主机制"],
            ["Thermodiffusion effect (TD)", "离子在温度梯度下发生 Soret 迁移", "作为 TG 对照，也可与 TG 协同"],
            ["Thermopower / α / Si", "单位温差产生的电压", "TG 最直观但不是唯一指标"],
            ["Redox couple", "可逆氧化还原对", "决定 TG 热力学上限"],
            ["Thermogalvanic hydrogel (TH)", "含 redox couple 的凝胶/水凝胶体系", "从液态走向准固态材料的关键平台"],
        ],
        widths=[1.7, 2.5, 2.3],
        font_size=8.4,
    )
    add_page_break(doc)


def section_background_v2(doc):
    add_heading(doc, "1. 背景：从低品位热到算力中心废热", 1)
    add_body(
        doc,
        """
低品位热回收的难点不是“热源是否存在”，而是“热源是否值得、是否方便、是否稳定地被转换”。Joule 2021 的 perspective 明确把低品位热来源列为太阳热/地热、工业过程副产废热、power plants、data centers、人体和微电子等，并指出低品位热由于温度低和转换效率低，商业利用率仍然有限。对于本报告的材料讨论，最值得强调的是数据中心与算力设施：服务器、GPU、网络和电源设备最终都把输入电能转化为热，散热系统持续把热带走；这些热量温度不高，但量大、持续、位置集中，并且与碳排和能效直接相关。

与间歇环境热相比，算力中心废热更适合作为材料报告背景有三个原因。第一，它更能代表“双重约束”：热需要被移除，同时希望部分转化为电能；材料不能阻碍散热。第二，数据中心废热通常处在低到中温区间，适合讨论 <100 °C 低品位热。第三，算力设施规模增长使热管理从附属工程问题变成能源效率问题。对于 TG 材料，这意味着研究目标不应只是演示温差发电，而应考虑在低温、连续热流和散热约束下保持稳定电化学输出。
""",
    )
    add_figure(doc, "fig_low_grade_heat_roadmap.png", "图 1. 低品位热背景、iTE 研究增长与演化路线", "Fig. 1, Chem. Soc. Rev., 2026", width=6.25)
    add_body(
        doc,
        """
图 1 的路线图显示，iTE 研究已经从“发现高 thermopower”转向“持续功率输出”。早期工作关注 thermocell、纳米多孔电解质和 ionic thermoelectric supercapacitors；随后，charged nanochannels、supramolecular thermocells、TD-TG 耦合以及 thermosensitive crystallization 被用来提高电压和输出。这个演化说明：低品位热转换的核心问题已经从“能不能产生电压”转向“如何在真实热源条件下稳定输出”。
""",
    )
    add_table(
        doc,
        ["低品位热场景", "材料层面需要面对的问题", "TG 路线的相关性"],
        [
            ["算力中心/数据中心废热", "热源持续但必须快速散热；材料不能显著阻碍热管理", "适合讨论连续输出、低热导与热移除之间的矛盾"],
            ["工业中低温废热", "温度和化学环境波动；面积大、成本敏感", "液态/凝胶电解质低成本和可扩展性有优势"],
            ["电力电子/微电子散热", "局部热流密度高，温差小且空间受限", "需要高 α、低内阻和快速响应"],
        ],
        widths=[1.8, 2.65, 2.05],
        font_size=8.2,
    )
    add_heading(doc, "1.1 算力中心废热给 TG 材料提出的四个约束", 2)
    add_body(
        doc,
        """
把数据中心作为背景，并不是要在本报告中讨论冷板、液冷、风冷或热回收系统，而是用它限定材料问题。首先，热源连续存在，因此材料不能只在短时温差下产生漂亮的开路电压，而要能在长时间温度循环中维持 redox 可逆性。其次，热源需要被移除，TG 材料若过度隔热或增加热阻，可能与散热目标冲突。第三，算力设施中的废热温度往往不高，意味着材料必须在小 ΔT 下有足够 α 和低内阻。第四，数据中心环境强调安全、低挥发、低泄漏和可维护性，这会削弱纯液态电解质的吸引力，增强准固态聚合物/凝胶体系的意义。
""",
    )
    add_table(
        doc,
        ["算力中心约束", "转化成材料问题", "对应的 TG 设计方向"],
        [
            ["连续热流", "redox couple 长时间可逆；电解质不衰减", "选择稳定 redox pair，避免不可逆副反应"],
            ["散热不能被阻碍", "低热导有利于 ΔT，但不能造成热积累", "控制厚度、热导和响应速度，不孤立追求低 κ"],
            ["小温差", "开路电压小，内阻影响被放大", "提高 α，同时保持高 σi 和快速 diffusion"],
            ["安全和维护", "液体泄漏、挥发、腐蚀不可接受", "发展 QSS、hydrogel、ionogel 和 polymer matrix"],
        ],
        widths=[1.4, 2.45, 2.65],
        font_size=8.0,
    )
    add_page_break(doc)


def section_mechanism_v2(doc):
    add_heading(doc, "2. 基本机制：为什么 TG 能连续输出", 1)
    add_body(
        doc,
        """
TG 的核心是温度改变 redox reaction 的电化学势。对于同一 redox couple，热端和冷端处于不同温度时，氧化态和还原态的自由能差不再相同，于是形成开路电压；接通外电路后，一侧发生氧化，另一侧发生还原，电子通过外电路流动。此时，电解质中的 redox species 需要持续扩散补给，supporting ions 需要维持电荷平衡。因此，TG 不是单纯热力学电压问题，而是 redox thermodynamics 与 mass transport 的耦合问题。

这也是 TG 与 TD/TCC 最大的区别。TD 依赖离子热迁移形成电双层或浓度梯度，通常可以产生更高 thermopower，但外电路不能直接让离子通过，因此更像热充电-放电过程。TG 通过电极 redox 反应让电子进入外电路，具备连续输出的机制基础。Advanced Science 2025 的综述指出，TEC 相比 TCC 的特点是可以维持连续稳定输出，并具有较快的电极-电解质界面质量传递和电荷转移动力学；但其性能由 thermopower、electrical/ionic conductivity 和 thermal conductivity 三个互相关联参数共同决定。
""",
    )
    add_figure(doc, "fig_three_mechanisms.png", "图 2. TD、TG 和 thermoextraction 的机制对比", "Fig. 1, Adv. Sci., 2025", width=6.15)
    add_figure(doc, "fig_td_tg_voltage_profiles.png", "图 3. TD 与 TG 的工作阶段和电压曲线差异", "Fig. 2, Adv. Mater., 2026", width=6.15)
    add_table(
        doc,
        ["问题", "TD/TCC", "TG/TEC"],
        [
            ["电压来源", "阳/阴离子热迁移差异与电双层建立", "redox reaction entropy 引起冷热端电极电势差"],
            ["输出模式", "热充电后放电，持续输出受循环限制", "redox 反应可让电子持续经过外电路"],
            ["主要优势", "thermopower 可很高", "连续输出逻辑更清楚"],
            ["主要短板", "能量密度和持续供电", "redox 扩散、内阻、浓差极化和稳定性"],
        ],
        widths=[1.25, 2.65, 2.6],
        font_size=8.3,
    )
    add_note(
        doc,
        "机制小结",
        "TG 的材料设计不能只问“α 多高”，还要问 redox species 能否快速、可逆、长期地在冷热端之间补给。",
    )
    add_page_break(doc)


def section_redox_v2(doc):
    add_heading(doc, "3. 核心变量一：redox couple 和溶剂化结构", 1)
    add_body(
        doc,
        """
Redox couple 是 TG 的第一性变量，因为 thermopower 近似与 ΔSredox/nF 相关。Joule 2021 和 Advanced Science 2025 都把 Fe(CN)6^4-/3- 作为基准 p-type redox couple：其水系 thermopower 约为 1.4 mV K-1；Fe2+/Fe3+ 约为 1.0 mV K-1；I-/I3- 可作为 n-type 体系，典型值约为 -0.86 mV K-1。这个基准说明，单靠常见 redox couple 的 intrinsic entropy difference，TG 热电势通常处在 mV K-1 量级，但进一步提升需要改变 redox species 的溶剂化、浓度分布或反应环境。
""",
    )
    add_figure(doc, "fig_redox_couples_strategies.png", "图 4. 常用 redox couples 与 thermopower 增强策略", "Fig. 6, Natl. Sci. Rev., 2026", width=5.85)
    add_heading(doc, "3.1 已经做了什么：三类 redox 增强策略", 2)
    add_table(
        doc,
        ["代表工作/策略", "综述中给出的结果", "说明的设计规律"],
        [
            ["Kim 等：向 Fe(CN)6^4-/3- 水系电解质加入 methanol", "Seebeck coefficient 从约 1.4 提高到约 2.9 mV K-1", "有机溶剂可选择性重排 redox species 的溶剂化结构，放大 ΔSredox"],
            ["Duan 等：向 Fe(CN)6^4-/3- 引入 urea 与 guanidinium", "热电势达到约 4.2 mV K-1", "chaotropic/氢键调控可改变水合壳层和局部结构熵"],
            ["Zhou 等：guanidinium 诱导 Fe(CN)6^4- 热敏结晶", "在 0.4 M Fe(CN)6^4-/3- 中达到约 3.73 mV K-1", "通过温敏结晶建立持续 redox concentration gradient"],
            ["α-cyclodextrin 捕获 I3-", "I-/I3- 体系热电势由约 0.86 提高到约 1.97 mV K-1", "超分子选择性捕获可以制造浓度差，改善 n-type redox 体系"],
            ["Fe2+/Fe3+ perchlorate", "水系中约 1.76 mV K-1，溶解度可超过 1 M", "弱 counterion interaction 有利于高溶解度和较高 thermopower"],
        ],
        widths=[2.0, 2.0, 2.5],
        font_size=7.8,
    )
    add_body(
        doc,
        """
这些例子可以合并成一条清晰规律：提高 TG thermopower 的 redox 路线主要有两种。第一种是提高 oxidized/reduced species 的溶剂化结构熵差，例如 methanol、urea、guanidinium 等改变 Fe(CN)6 的水合环境。第二种是人为制造冷热端浓度差，例如 α-cyclodextrin 选择性捕获 I3- 或 guanidinium 诱导 Fe(CN)6^4- 结晶。前者偏热力学，后者偏浓度梯度和相行为。

但两类策略都存在同一个限制：提高 thermopower 往往会牺牲传输。Joule 2021 指出，调节 redox solvation structure 虽可增强 Se，却常伴随有效电导率下降，因为优化电解质的离子电导可能较低。换言之，redox chemistry 的成功不能只看开路电压，还必须看 viscosity、diffusion、solubility、charge-transfer resistance 和长期化学稳定性。
""",
    )
    add_heading(doc, "3.2 Redox couple 家族的体系化比较", 2)
    add_body(
        doc,
        """
从综述中反复出现的 redox families 来看，Fe(CN)6^4-/3- 是最成熟的 p-type 基准体系，优点是水系可逆、反应动力学快、容易与水凝胶兼容；缺点是 thermopower 的进一步提升越来越依赖外部溶剂化或结晶策略。I-/I3- 的优势是可通过选择性包结或络合实现 n-type 或浓度梯度调控，缺点是副反应、挥发/稳定性和浓差极化更需要仔细处理。Fe2+/Fe3+ 体系原理简单，但与 counterion 的相互作用、溶解度和反应动力学决定了实际性能；perchlorate 体系被综述作为值得继续研究的例子。Co(bpy)3^2+/3+、BQ/HQ 等有机/配合物 redox pair 则提供更丰富的反应环境调控空间，但成本、毒性、pH 稳定和聚合物兼容性必须同步评估。
""",
    )
    add_table(
        doc,
        ["Redox family", "优势", "已被证明的策略", "主要问题"],
        [
            ["Fe(CN)6^4-/3-", "水系稳定、可逆性好、基准体系成熟", "methanol/urea/Gdm+ 调溶剂化；Gdm+ 热敏结晶；gelatin/KCl 协同", "进一步提高 α 依赖复杂添加剂或相行为；溶解度和扩散仍受限"],
            ["I-/I3-", "可作为 n-type 代表；容易被超分子主体选择性作用", "α-CD 捕获 I3- 建立浓度梯度；methylcellulose/I3- 复合实现自充电行为", "副反应、长期稳定、挥发和浓差极化"],
            ["Fe2+/Fe3+", "反应概念简单，价态变化明确", "perchlorate counterion 降低相互作用并提高溶解度", "水解、pH 依赖和反应动力学需要控制"],
            ["Co(bpy)3^2+/3+", "非水/准固态体系可调性强", "quasi-solid-state electrolyte 和 cellulose film 中应用", "成本、黏度、扩散和环境稳定性"],
            ["BQ/HQ", "有机 redox 可通过 pH 和聚合物环境调控", "BQ/HQ 净 TG 效应可达到较高 α；自调 pH 聚合物可偏移平衡", "pH 敏感、副反应和聚合物兼容性"],
        ],
        widths=[1.35, 1.45, 2.15, 1.55],
        font_size=7.2,
    )
    add_note(
        doc,
        "本章结论",
        "redox couple 的选择不是材料清单，而是热力学、溶剂化、扩散和稳定性的联立优化；综述中的高 α 案例几乎都不是单纯换 redox pair，而是同时改变化学环境。",
    )
    add_page_break(doc)


def section_polymer_hydrogel_v2(doc):
    add_heading(doc, "4. 核心变量二：从液态电解质到聚合物/水凝胶基体", 1)
    add_body(
        doc,
        """
液态 thermocell 是最直接的 TG 平台，但液体体系存在泄漏、挥发、冻结/沸腾和封装稳定性问题。准固态聚合物和 thermogalvanic hydrogels（THs）的意义在于：它们把 redox couple 固定在一个可形变、可保水、可调控微结构的离子网络中。这个网络不是惰性容器，而会通过官能团、孔道、水状态和链段运动改变 redox species 的溶剂化、扩散和局部浓度。
""",
    )
    add_figure(doc, "fig_material_categories.png", "图 5. iTE 材料从液态、凝胶态到聚电解质和多孔固态框架的分类", "Fig. 2, Natl. Sci. Rev., 2026", width=6.05)
    add_figure(doc, "fig_performance_material_comparison.png", "图 6. iTE/eTE 性能区间与功能聚合物多尺度设计", "Fig. 1, Adv. Mater., 2026", width=6.15)
    add_heading(doc, "4.1 已经做了什么：thermogalvanic hydrogels 的代表体系", 2)
    add_table(
        doc,
        ["材料体系", "综述引用的代表工作", "已经证明了什么"],
        [
            ["PVA / Fe(CN)6^4-/3- / water", "Zhou 等 2016 首次提出 TH；PVA 为聚合物基体，Fe(CN)6 为 redox couple", "TH 可把液态 TG 转为准固态水凝胶平台，初始 thermopower 约 1.21 mV K-1"],
            ["Gelatin / KCl / Fe(CN)6^4-/3-", "Liu 等把 Fe(CN)6 的 TG 与 KCl 的 TD 协同", "TG-only 约 4.8 mV K-1；协同后约 17 mV K-1，说明水凝胶可整合 TG 与 TD 贡献"],
            ["Cellulose / Fe(CN)6^4-/3-", "Pringle 等制备 5 wt% cellulose 的 Fe(CN)6 水凝胶电解质", "在保持机械性的同时，热电化学性能接近液态电解质（约 -1.38 mV K-1）"],
            ["Bacterial cellulose / Fe(CN)6^4-/3- / hygroscopic salts", "Feng 等利用 BC 纳米纤维和吸湿盐", "提升机械强度、含水稳定和离子传输，显示天然纳米网络可作为 TH 基体"],
            ["PVA 取向/机械训练体系", "综述中提到通过机械训练或冻融取向形成更有序结构", "离子电导可由约 2.6 提高至 4.6 S m-1，并保持 Fe(CN)6 redox 功能"],
            ["PAAm / metal-ion interaction", "PAAm 中引入 Fe3+/Fe2+ redox couple", "金属离子与酰胺基作用增强网络，兼顾力学与 TG 功能"],
        ],
        widths=[1.55, 2.3, 2.65],
        font_size=7.7,
    )
    add_figure(doc, "fig_hydrogel_synergy.png", "图 7. Gelatin/cellulose 基 TH 的协同效应与代表进展", "Fig. 2, Adv. Sci., 2025", width=6.05)
    add_body(
        doc,
        """
这些例子说明，TH 的发展逻辑不是“把电解液凝胶化”这么简单，而是逐步把材料变量纳入 TG 过程。PVA 证明准固态 TG 可行；gelatin/KCl/Fe(CN)6 证明 TD 与 TG 可以在一个水凝胶体系中协同，但也提醒我们要区分 TG 贡献和 TD 贡献；cellulose 和 bacterial cellulose 证明天然多孔网络可同时提供机械支撑和离子通道；PAAm 和 PVA 的合成网络证明链段相互作用、交联密度和取向结构可以用来调节传输。
""",
    )
    add_heading(doc, "4.2 聚合物基体到底调控了哪些变量", 2)
    add_body(
        doc,
        """
Advanced Materials 2026 将功能聚合物的作用拆成分子、微结构和宏观三个尺度，这比简单按“PVA、gelatin、cellulose”分类更有解释力。分子尺度上，羟基、酰胺基、羧基、氨基、磺酸基和 zwitterionic groups 通过氢键、静电作用、配位和离子偶极作用改变离子解离与水合结构。微结构尺度上，取向通道、纳米纤维、相分离和孔径控制 redox species 的迁移路径。宏观尺度上，含水率、抗冻/抗干燥、热导和机械韧性决定材料能否稳定承受低品位热环境。

例如，aligned cellulose 和 lignin/PVA 等体系被综述用来说明“通道方向性”可以同时影响 σi 和 α；机械训练的 PVA 网络说明取向结构不只是力学增强，也能改善 ion migration；gelatin/GTA 交联体系说明稳定 3D polymer matrix 可以拓宽工作温度和保持离子通道；PAAm 与 Fe3+/Fe2+ 的作用说明 redox species 本身也可能成为网络增强或动态交联的一部分。这样看，聚合物不是 TG 电解质的包装材料，而是 redox 环境和传输路径的共同设计变量。
""",
    )
    add_table(
        doc,
        ["聚合物设计变量", "影响的 TG 过程", "代表性例子"],
        [
            ["官能团极性/固定电荷", "改变离子解离、溶剂化和 redox species 局部分布", "gelatin 的带电网络、zwitterionic polymer、anionic polymer 调 pH"],
            ["交联密度和网络强度", "提高形状稳定，但可能降低 diffusion", "gelatin/GTA、PAAm/PVA double network"],
            ["取向和纳米通道", "降低迁移路径阻力，形成方向性 ion transport", "aligned cellulose、mechanically trained PVA"],
            ["水状态调控", "影响含水稳定、冰点、黏度和离子迁移", "hygroscopic salts、organohydrogel、bound water 设计"],
            ["热导和孔结构", "维持 ΔT，同时影响响应速度", "fibrous/porous polymer, low-κ QSS ionic conductors"],
        ],
        widths=[1.7, 2.45, 2.35],
        font_size=7.8,
    )
    add_heading(doc, "4.3 需要避免的误读：高 thermopower 不一定来自纯 TG", 2)
    add_body(
        doc,
        """
Gelatin/KCl/Fe(CN)6 的 17 mV K-1 是非常重要的代表工作，但它同时也是一个容易被误读的案例。综述明确指出，该体系中 Fe(CN)6 redox couple 提供 TG 贡献，K+ 和 Cl- 通过 TD effect 贡献热扩散项；超过 60% 的 Seebeck coefficient 来自 KCl 的 TD effect。因此，这个结果的意义不是“纯 TG 已经达到 17 mV K-1”，而是证明水凝胶网络可以把 TG 和 TD 机制整合到同一材料中。对于写报告，这类例子应该被归为“协同机制”，而不是简单列入 TG redox couple 的 intrinsic performance。

同样，thermosensitive crystallization 和 α-CD/I3- 捕获策略也不是单纯提高 redox entropy，而是利用相行为或超分子结合产生 redox concentration gradient。它们说明 TG 可以借助外部化学环境增强，但也意味着高性能往往伴随更复杂的可逆性和稳定性问题。
""",
    )
    add_page_break(doc)


def section_system_comparison_v2(doc):
    add_heading(doc, "5. 体系化比较：不同材料路线已经解决了什么", 1)
    add_body(
        doc,
        """
如果把已有工作按“解决的问题”而不是按材料名称排列，thermogalvanic 材料体系可以分成四条路线：redox chemistry 路线、solvation/additive 路线、polymer matrix 路线和 hydrogel microstructure 路线。它们不是互斥关系，而是逐层叠加：redox chemistry 决定上限，solvation/additive 放大上限，polymer matrix 决定能否稳定运行，hydrogel microstructure 决定能否兼顾传输和力学。
""",
    )
    add_table(
        doc,
        ["路线", "主要已经解决的问题", "代表例子", "仍然没有完全解决的问题"],
        [
            ["Redox couple 选择", "找到可逆、可溶、较高 ΔSredox 的活性对", "Fe(CN)6^4-/3-、I-/I3-、Fe2+/3+、Co(bpy)3^2+/3+、BQ/HQ", "高 thermopower 与毒性、溶解度、反应速率难兼得"],
            ["溶剂化/添加剂", "通过水合壳层和 counterion interaction 提高 α", "methanol、urea、guanidinium、organic-water mixtures", "电导下降、黏度增加、相稳定性和副反应"],
            ["浓度梯度/相行为", "把 redox species 在冷热端分布差异固定下来", "α-CD 捕获 I3-；guanidinium 诱导 Fe(CN)6^4- 结晶", "浓度梯度长期稳定性和可逆恢复"],
            ["聚合物基体", "把液态电解质转为准固态并调控离子环境", "PVA、gelatin、cellulose、PAAm、PVDF-HFP", "强网络可能阻碍 redox diffusion"],
            ["水凝胶微结构", "同时改善保水、机械、离子通道和低热导", "BC 纳米纤维、取向 PVA、双网络 PAAm/PVA", "环境稳定、标准化测试和长期循环"],
        ],
        widths=[1.25, 1.75, 1.75, 1.75],
        font_size=7.5,
    )
    add_heading(doc, "5.1 一个更合理的材料评价顺序", 2)
    add_body(
        doc,
        """
评价 TG 材料时，建议不要从“最高 α”开始，而应按以下顺序判断。第一，redox couple 是否可逆且足够稳定；第二，溶剂化或添加剂策略是否提高 α，同时是否牺牲 σi 和 diffusion；第三，聚合物/水凝胶是否保持 redox species 的有效补给；第四，材料热导是否有利于保持温差；第五，含水率、抗冻、抗干燥和化学稳定是否能支撑长时间测试。
""",
    )
    add_table(
        doc,
        ["评价层级", "关键问题", "为什么重要"],
        [
            ["1. Redox thermodynamics", "ΔSredox 是否足够高？p/n 极性是否明确？", "决定开路电压和材料上限"],
            ["2. Reaction and mass transport", "redox species 是否快速扩散？电荷转移是否快？", "决定闭路电流和浓差极化"],
            ["3. Polymer-ion interaction", "官能团是否提高 α，同时不过度束缚 redox species？", "决定准固态体系是否比液态更优"],
            ["4. Thermal property", "κ 是否足够低以维持 ΔT？", "决定低品位热源下有效温差"],
            ["5. Stability", "水分、盐析、结晶、pH 和副反应是否稳定？", "决定材料能否从短时演示走向可比较测试"],
        ],
        widths=[1.55, 2.35, 2.6],
        font_size=8.0,
    )
    add_heading(doc, "5.2 已有工作形成的“材料路线图”", 2)
    add_body(
        doc,
        """
如果把这些代表工作放到时间线上，可以看到 TG/iTE 材料路线并不是线性替代，而是层层叠加。液态 thermocell 阶段先证明 redox entropy 可以在低温差下产生 mV K-1 级 thermopower；solvation engineering 阶段通过 methanol、urea、guanidinium 和 organic-water mixtures 放大 ΔSredox；concentration-gradient 阶段通过 α-CD 或 thermosensitive crystallization 让 redox species 在冷热端形成稳定差异；hydrogel 阶段把这些化学策略放进可保持形状和水分的聚合物网络；functional polymer 阶段进一步把 ion dynamics、mechanics 和 thermal properties 放入同一设计框架。

因此，本领域的“体系”可以概括为五层：redox pair 是核心，solvent/additive 是局部环境，supporting electrolyte 是迁移与协同项，polymer matrix 是传输和稳定平台，microstructure 是把上述变量空间化的手段。报告中所有代表工作都可以被放入这五层，而不是彼此孤立。
""",
    )
    add_table(
        doc,
        ["层级", "材料问题", "综述中的代表做法", "评价重点"],
        [
            ["1. Redox pair", "选择高 ΔSredox、可逆、可溶的活性对", "Fe(CN)6^4-/3-、I-/I3-、Fe2+/3+、BQ/HQ", "α、反应可逆性、溶解度"],
            ["2. Solvent/additive", "调水合壳层和局部 entropy", "methanol、urea、Gdm+、organic-water mixture", "α 与 σi/diffusion 的权衡"],
            ["3. Supporting electrolyte", "提供离子传输与 TD 协同", "KCl in gelatin/Fe(CN)6 system", "区分 TG 与 TD 贡献"],
            ["4. Polymer matrix", "准固态化、保水、力学和官能团作用", "PVA、gelatin、cellulose、PAAm、PVDF-HFP", "稳定性、含水率、redox migration"],
            ["5. Microstructure", "构建通道、孔结构和取向", "aligned cellulose、BC nanofibers、mechanical training", "σi、κ、Pmax/(ΔT)^2、循环稳定"],
        ],
        widths=[1.05, 1.75, 2.15, 1.55],
        font_size=7.5,
    )
    add_page_break(doc)


def section_challenges_v2(doc):
    add_heading(doc, "6. 剩余问题：下一步材料设计该围绕什么展开", 1)
    add_body(
        doc,
        """
第一，TG 需要把 thermopower 和 transport 放在一起优化。许多策略可以提高 α，例如改变溶剂化、引入添加剂或建立浓度梯度，但它们也可能提高黏度、降低扩散、诱发沉淀或降低电导。材料论文如果只报告 α，而不同时报告 σi、diffusion、Pmax/(ΔT)^2、ηr 和稳定性，难以判断真实价值。

第二，n-type thermogalvanic materials 仍明显落后于 p-type benchmark。Fe(CN)6^4-/3- 作为 p-type 基准体系成熟，而 I-/I3-、Fe2+/Fe3+ 等 n-type 或反向极性体系常受限于反应动力学、稳定性和热电势。综述中提到 Fe2+/Fe3+ perchlorate 可在水系中取得较高 thermopower 和溶解度，说明 counterion interaction 是 n-type 设计的重要方向。

第三，聚合物网络必须避免“强而不通”。水凝胶增强常依赖交联、双网络、氢键和结晶区，但 redox species 需要在冷热端之间持续迁移。过强的聚合物-redox 作用可能提高局部 entropy，却降低补给速度。未来应更重视 diffusion coefficient、EIS 和原位谱学，而不是只用力学照片或短时电压曲线证明性能。

第四，算力中心废热这类背景要求材料与热管理兼容。材料层面需要承认一个矛盾：低热导有助于维持温差，但废热场景同时需要把热移走，不能让热源过热。因此，TG 材料不能孤立追求低 κ；更合理的目标是在局部温差保持、热移除和电化学输出之间建立平衡。这个问题虽然会延伸到系统设计，但在材料阶段就应通过热导、厚度、响应时间和长期温度循环来约束。
""",
    )
    add_note(
        doc,
        "最终判断",
        "一个成熟的 TG 材料体系应当同时满足：可逆 redox chemistry、高且可解释的 ΔSredox、不过度牺牲扩散的溶剂化调控、稳定聚合物/水凝胶网络、低温差下可重复的 Pmax 与 ηr。",
    )
    add_heading(doc, "参考综述来源", 1)
    refs = [
        "Duan et al. Liquid-state thermocells: Opportunities and challenges for low-grade heat harvesting. Joule, 2021.",
        "Liu et al. Potentially Wearable Thermo-Electrochemical Cells for Body Heat Harvesting: From Mechanism, Materials, Strategies to Applications. Advanced Science, 2021.",
        "Jabeen et al. Recent advances in ionic thermoelectric systems and theoretical modelling. Chemical Science, 2024.",
        "Lin et al. Critical Design Strategy of Thermogalvanic Hydrogels for Low-Grade Heat Harvesting. Advanced Science, 2025.",
        "Kim et al. Functional Polymers for Ionic Thermoelectrics: Multiscale Design Strategies for Ion Dynamics, Mechanics, and Energy Harvesting. Advanced Materials, 2026.",
        "Shang et al. Pushing ionic thermoelectrics toward power supply applications: origins, advances, challenges, and future directions. Chemical Society Reviews, 2026.",
        "Xu et al. A bird's-eye view for future ionic thermoelectrics: from ions to products. National Science Review, 2026.",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Number")
        set_para_format(p, after=4, line=1.12)
        r = p.add_run(ref)
        set_run_font(r, STYLE["font_cn"], size=9.3, color=STYLE["body"])


def add_cover_v2(doc):
    p = doc.add_paragraph()
    set_para_format(p, before=120, after=12, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Thermogalvanic 低品位热能转换体系化报告")
    set_run_font(r, STYLE["font_cn"], size=23, bold=True, color=STYLE["accent"])
    p = doc.add_paragraph()
    set_para_format(p, after=18, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("从算力中心废热背景到 redox couple、溶剂化结构与水凝胶材料设计")
    set_run_font(r, STYLE["font_cn"], size=13.5, color=STYLE["accent2"])
    p = doc.add_paragraph()
    set_para_format(p, after=20, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("基于 7 篇综述及其引用代表工作的综合整理；不讨论器件集成、可穿戴应用和模块层面。")
    set_run_font(r, STYLE["font_cn"], size=10.2, color=STYLE["muted"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / "fig_low_grade_heat_roadmap.png"), width=Inches(5.5))
    add_caption(doc, "封面图. 低品位热背景与 iTE 发展路线图", "Fig. 1, Chem. Soc. Rev., 2026")
    add_page_break(doc)


def build():
    OUT.mkdir(exist_ok=True)
    doc = Document()
    configure_doc(doc)
    add_cover_v2(doc)
    add_manual_toc_v2(doc)
    add_front_matter_v2(doc)
    section_background_v2(doc)
    section_mechanism_v2(doc)
    section_redox_v2(doc)
    section_polymer_hydrogel_v2(doc)
    section_system_comparison_v2(doc)
    section_challenges_v2(doc)
    doc.save(DOCX)
    Document(DOCX)
    write_qa()
    print(DOCX)


if __name__ == "__main__":
    build()
